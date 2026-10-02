"""
CardioTwin-India: Architecture Diagram Generator
Generates:
1. docs/architecture/architecture_diagram.png (High-res 300 DPI)
2. docs/architecture/CardioTwin_Architecture_Diagram.pdf
3. docs/architecture/CardioTwin_Architecture_Diagram.pptx
"""

import os
import matplotlib.pyplot as plt
import matplotlib.patches as patches
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def generate_diagram_image(output_png):
    os.makedirs(os.path.dirname(output_png), exist_ok=True)
    fig, ax = plt.subplots(figsize=(16, 10), dpi=300)
    ax.set_xlim(0, 16)
    ax.set_ylim(0, 10)
    ax.axis("off")
    fig.patch.set_facecolor("#0F172A") # Modern Dark Tech Slate

    # Title
    ax.text(8, 9.5, "CardioTwin-India: Multimodal Digital Twin System Architecture",
            fontsize=20, fontweight="bold", color="#F8FAFC", ha="center")
    ax.text(8, 9.15, "Official Architecture Blueprint for Happiest Health Summit 2026",
            fontsize=12, color="#94A3B8", ha="center")

    def draw_layer_box(x, y, w, h, bg_color, border_color, title, subtitle):
        rect = patches.FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.15",
                                      facecolor=bg_color, edgecolor=border_color, linewidth=2)
        ax.add_patch(rect)
        ax.text(x + 0.3, y + h - 0.35, title, fontsize=11, fontweight="bold", color="#F8FAFC")
        ax.text(x + 0.3, y + h - 0.65, subtitle, fontsize=8.5, color="#CBD5E1")

    def draw_component(x, y, w, h, bg_color, border_color, text, subtext=""):
        rect = patches.FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.08",
                                      facecolor=bg_color, edgecolor=border_color, linewidth=1.2)
        ax.add_patch(rect)
        ax.text(x + w/2, y + h/2 + (0.1 if subtext else 0), text, fontsize=9.5, fontweight="bold",
                color="#FFFFFF", ha="center", va="center")
        if subtext:
            ax.text(x + w/2, y + h/2 - 0.22, subtext, fontsize=7.5, color="#E2E8F0", ha="center", va="center")

    # Layer 1: Data Ingestion (Bottom)
    draw_layer_box(0.8, 0.6, 14.4, 1.6, "#1E293B", "#38BDF8",
                   "LAYER 1: MULTIMODAL PATIENT INGESTION & SENSING",
                   "Longitudinal Electronic Health Records (EHR) & High-Frequency Wearable IoT Telemetry Streams")
    
    draw_component(1.2, 0.8, 3.2, 0.9, "#0369A1", "#38BDF8", "Static Clinical EHR", "Labs, HbA1c, Lipids, UACR, Meds")
    draw_component(4.7, 0.8, 3.2, 0.9, "#0D9488", "#2DD4BF", "Continuous Glucose (CGM)", "5-min interstitial glucose flux")
    draw_component(8.2, 0.8, 3.2, 0.9, "#E11D48", "#FB7185", "Wearable PPG Sensor", "Continuous Pulse, HRV (RMSSD)")
    draw_component(11.7, 0.8, 3.2, 0.9, "#7C3AED", "#A78BFA", "Continuous BP & Activity", "Ambulatory SBP/DBP & Step METs")

    # Connector 1 -> 2
    for cx in [2.8, 6.3, 9.8, 13.3]:
        ax.annotate("", xy=(cx, 2.5), xytext=(cx, 2.2),
                    arrowprops=dict(arrowstyle="->", color="#38BDF8", lw=2))

    # Layer 2: Preprocessing & Fusion
    draw_layer_box(0.8, 2.5, 14.4, 1.4, "#1E293B", "#818CF8",
                   "LAYER 2: TEMPORAL DATA FUSION & PHYSIOLOGICAL ALIGNMENT",
                   "Time-series synchronization, autoregressive imputation, and Asian Indian phenotypic normalization")
    
    draw_component(1.5, 2.7, 3.8, 0.8, "#312E81", "#818CF8", "Multivariate Imputation", "Kalman Filtering & Spline Interpolation")
    draw_component(6.1, 2.7, 3.8, 0.8, "#312E81", "#818CF8", "Circadian Normalizer", "Diurnal Rhythm & Fasting-State Sync")
    draw_component(10.7, 2.7, 3.8, 0.8, "#312E81", "#818CF8", "Feature Engineering", "TIR (70-180), Glycemic Lability, RPP")

    # Connector 2 -> 3
    for cx in [3.4, 8.0, 12.6]:
        ax.annotate("", xy=(cx, 4.2), xytext=(cx, 3.9),
                    arrowprops=dict(arrowstyle="->", color="#818CF8", lw=2))

    # Layer 3: Physiological Digital Twin Core
    draw_layer_box(0.8, 4.2, 14.4, 2.3, "#1E293B", "#C084FC",
                   "LAYER 3: CORE PHYSIOLOGICAL DIGITAL TWIN & AI PREDICTIVE ENGINE",
                   "Coupled Bio-Physical Compartmental Models & Multi-Horizon Machine Learning Ensembles")
    
    draw_component(1.2, 4.5, 3.2, 1.5, "#4C1D95", "#C084FC", "Bergman Minimal Model\n(ODE System)",
                   "dG/dt, dX/dt, dI/dt Dynamics\nInsulin Sensitivity & Beta-Cell")
    draw_component(4.7, 4.5, 3.2, 1.5, "#4C1D95", "#C084FC", "2-Element Windkessel\nHemodynamics",
                   "dP/dt Arterial Resistance & Comp.\nNocturnal BP Dipping Profile")
    draw_component(8.2, 4.5, 3.2, 1.5, "#065F46", "#34D399", "120-min Seq Forecaster\n(AR / Recurrent)",
                   "Short-Term Glucose Horizon\nImpending Hypoglycemia Alerts")
    draw_component(11.7, 4.5, 3.2, 1.5, "#991B1B", "#F87171", "XGBoost Multi-Target\nClinical Ensembles",
                   "10-Yr ASCVD (1.5x Indian Mult.)\nDKD Stage & 90-d Decompensation")

    # Connector 3 -> 4
    for cx in [2.8, 6.3, 9.8, 13.3]:
        ax.annotate("", xy=(cx, 6.8), xytext=(cx, 6.5),
                    arrowprops=dict(arrowstyle="->", color="#C084FC", lw=2))

    # Layer 4: Application & Clinical Decision Support
    draw_layer_box(0.8, 6.8, 14.4, 2.0, "#1E293B", "#34D399",
                   "LAYER 4: CLINICAL INTERACTION, COUNTERFACTUAL SIMULATION & INTEROPERABILITY",
                   "Streamlit Doctor/Patient Interface, Fast In-Silico Testing, Explainability, and HL7 FHIR Integration")
    
    draw_component(1.2, 7.1, 3.2, 1.2, "#064E3B", "#34D399", "Live Twin Dashboard", "Multi-stream telemetry,\norgan status radar gauges")
    draw_component(4.7, 7.1, 3.2, 1.2, "#064E3B", "#34D399", "What-If Simulator", "Test low-carb meals, walks,\nSGLT2i & ARB titration")
    draw_component(8.2, 7.1, 3.2, 1.2, "#1E3A8A", "#60A5FA", "Explainability Engine", "ICMR-INDIAB phenotype rules,\nSHAP driver attributions")
    draw_component(11.7, 7.1, 3.2, 1.2, "#581C87", "#D8B4FE", "HL7 FHIR & ABDM API", "FastAPI microservice,\nPatient & Observation Bundles")

    plt.tight_layout()
    plt.savefig(output_png, facecolor=fig.get_facecolor(), edgecolor="none")
    plt.close()
    print(f"Generated High-Res Architecture PNG at {output_png}")

def generate_diagram_pptx(output_pptx, diagram_png):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_layout = prs.slide_layouts[6] # Blank
    slide = prs.slides.add_slide(slide_layout)

    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(15, 23, 42) # Slate 900
    bg.line.fill.background()

    # Title
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.0))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "CardioTwin-India: Architecture Blueprint"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(248, 250, 252)

    p2 = tf.add_paragraph()
    p2.text = "Multimodal Physiomarker Digital Twin for Type-2 Diabetes & Cardio-Metabolic Progression in India"
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(148, 163, 184)

    # Add diagram image
    slide.shapes.add_picture(diagram_png, Inches(0.8), Inches(1.6), width=Inches(11.733))

    prs.save(output_pptx)
    print(f"Generated Architecture PPTX at {output_pptx}")

def generate_diagram_pdf(output_pdf, diagram_png):
    doc = SimpleDocTemplate(
        output_pdf,
        pagesize=landscape(letter),
        leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=22,
        leading=26,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=6
    )
    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontSize=11,
        textColor=colors.HexColor('#475569'),
        spaceAfter=14
    )

    story = []
    story.append(Paragraph("CardioTwin-India: System Architecture Blueprint", title_style))
    story.append(Paragraph("Happiest Health - Digital Twin Challenge 2026 | Reimagining & Reforming Healthcare in India Summit", sub_style))
    story.append(RLImage(diagram_png, width=720, height=450))
    doc.build(story)
    print(f"Generated Architecture PDF at {output_pdf}")

def main():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    arch_dir = os.path.join(root, "docs", "architecture")
    os.makedirs(arch_dir, exist_ok=True)
    
    png_path = os.path.join(arch_dir, "architecture_diagram.png")
    pdf_path = os.path.join(arch_dir, "CardioTwin_Architecture_Diagram.pdf")
    pptx_path = os.path.join(arch_dir, "CardioTwin_Architecture_Diagram.pptx")
    
    generate_diagram_image(png_path)
    generate_diagram_pptx(pptx_path, png_path)
    generate_diagram_pdf(pdf_path, png_path)

if __name__ == "__main__":
    main()
